import argparse
import sqlite3
import datetime
import os
from pathlib import Path
from openpyxl import load_workbook
from fpdf import FPDF

db = None
debug = True

def adapt_date_iso(val):
    return val.isoformat()
  
def main():
  parser = argparse.ArgumentParser(
                    prog='Classifiche',
                    description='Classifiche CrazyScalata')
  parser.add_argument('filename')
  args = parser.parse_args()
  setup(args.filename)
  create_rankings()
  create_extra()
  db.close()
  
def setup(fname: str):
  sqlite3.register_adapter(datetime.datetime, adapt_date_iso)
  create_db()
  read_data(fname)
  
def read_data(fname: str):
  p = Path(fname)
  wb = load_workbook(filename = p.absolute())
  wb.active = wb["ISCRITTI"]
  row = 2
  ws = wb.active
  cursor = db.cursor()
  while (True):
    if ws.cell(row=row, column=1).value is None:
      break
    cursor.execute("INSERT INTO iscritti VALUES(?, ?, ?, ?, ?, ?, ?, ?)", [row] + list(map(lambda x : x.value, ws[row])))
    row = row + 1
  db.commit()
  
  wb.active = wb["TEMPI"]
  row = 2
  ws = wb.active
  cursor = db.cursor()
  while (True):
    if ws.cell(row=row, column=1).value is None:
      break
    cursor.execute("INSERT INTO tempi VALUES(?, ?, ?, ?)", list(map(lambda x : x.value, ws[row])) + process_time(ws[row][1].value))
    row = row + 1
  db.commit()
  
def process_time(t: str) -> list[any]:
  if "." not in t and "," not in t:
    # no fractional part, add it
    t = t + ".00"
    
  if "," in t:
    t = t.replace(",", ".")
    
  whole, frac = t.split(".", 1)
  frac = frac.ljust(2, "0")[0:2]
  
  wparts = whole.split(":")
  ftime = float(frac) / 100  
  
  if len(wparts) == 1:
    secs = wparts[0].rjust(2, "0")
    ftime += float(secs)
    whole = "00:00:" + secs
  if len(wparts) == 2:
    secs = wparts[1].rjust(2, "0")
    mins = wparts[0].rjust(2, "0")
    ftime += float(secs) + 60 * float(mins)
    whole = "00:" + mins + ":" + secs
  if len(wparts) == 3:
    secs = wparts[2].rjust(2, "0")
    mins = wparts[1].rjust(2, "0")
    hours = wparts[0].rjust(2, "0")
    ftime += float(secs) + 60 * float(mins) + 60 * 60 * float(hours)
    whole = hours + ":" + mins + ":" + secs
    
  res = [whole + "." + frac, ftime]
  return res

def format_date(dt: str):
  isodt = dt.split("T")[0].split("-")
  isodt.reverse()
  return "/".join(isodt)
  
def extract_ranking_data(table: str) -> list[list[str]]:
  data = []
  hdr = ['Pos.', 'Nome', 'Pettorale', 'Tempo', 'Data Nascita', 'Crazy']
  c = db.cursor()
  c.execute(f"SELECT * FROM {table}")
  rows = c.fetchall()
  for idx, r in enumerate(rows):
    data.append([str(idx + 1), r[0], str(r[1]), r[6], format_date(r[4]), r[3]])
  c.close()
  return [hdr] + data
  
def create_rankings():
  create_ranking("Assoluti", extract_ranking_data("c_assoluti"))
  create_ranking("Maschile", extract_ranking_data("c_uomini"))
  create_ranking("Femminile", extract_ranking_data("c_donne"))
  create_ranking("CrazyBikers", extract_ranking_data("c_crazy"))
  create_ranking("Ciclisti", extract_ranking_data("c_ciclisti"))
  create_ranking("Podisti", extract_ranking_data("c_podisti"))
  create_ranking("Data Nascita", extract_ranking_data("c_eta"))
  
def format_time(t: float) -> str:
  h_secs = 60 * 60
  m_secs = 60
  
  h = int(t / h_secs)
  remainder = t - h * h_secs
  m = int(remainder / m_secs)
  remainder = remainder - m * m_secs
  s = int(remainder)
  frac = remainder - s
  
  wparts = []
  if h > 0: wparts.append(str(h).rjust(2, "0"))
  if m > 0: wparts.append(str(m).rjust(2, "0"))
  wparts.append(str(s).rjust(2, "0"))
  
  whole = ":".join(wparts)
  if frac > 0:
    res = whole + str(frac).ljust(2, "0")[1:4]
    
  return res
    
  
def create_extra():
  c = db.cursor()
  c.execute("SELECT AVG(ftempo) FROM c_crazy WHERE tipo = 'CICLISTA'")
  avg = c.fetchone()[0]
  c.execute(f"SELECT *, abs({avg} - ftempo) as diff FROM c_crazy WHERE tipo = 'CICLISTA' ORDER BY diff ASC LIMIT 10")

  data = [['Pos.', 'Nome', 'Pettorale', 'Tempo', 'Data Nascita', 'Crazy']]
  rows = c.fetchall()
  for idx, r in enumerate(rows):
    data.append([str(idx + 1), r[0], str(r[1]), r[6], format_date(r[4]), r[3]])
  c.close()
  
  pdf = FPDF(orientation="landscape", format="A4")
  pdf.add_page()
  pdf.set_font("helvetica", style="B", size=16)
  pdf.set_font_size(18)
  pdf.write(text="CrazyScalata 2025\n")
  pdf.set_font("helvetica", style="", size=12)
  pdf.write(10, "I più vicini alla media\n")
  pdf.write(10, f'Media Ciclisti Crazy: {format_time(avg)}\n')
  
  spacing = 1.15
  row_height = pdf.font_size
  colw = [20, 100, 40, 40, 50, 20]
  for row in data:
    for idx, item in enumerate(row):
      pdf.cell(colw[idx], row_height*spacing, text=item, border=1)
    pdf.ln(row_height*spacing)
  pdf.output("out/Media.pdf")

def create_ranking(name: str, data: list[list[str]]):
  pdf = FPDF(orientation="landscape", format="A4")
  pdf.add_page()
  pdf.set_font("helvetica", style="B", size=16)
  pdf.set_font_size(18)
  pdf.write(text="CrazyScalata 2025\n")
  pdf.set_font("helvetica", style="", size=12)
  pdf.write(10, "Classifica " + name)
  pdf.write(text='\n\n')
  spacing = 1.15
  row_height = pdf.font_size
  colw = [20, 100, 40, 40, 50, 20]
  for row in data:
    for idx, item in enumerate(row):
      pdf.cell(colw[idx], row_height*spacing, text=item, border=1)
    pdf.ln(row_height*spacing)
  pdf.output("out/{0}.pdf".format(name))
  
def create_db():
  global db
  global debug
  
  if (debug):
    debug_db = "debug.sqlite"
    os.remove(debug_db)
    db = sqlite3.connect(debug_db)
  else:
    db = sqlite3.connect(":memory:")
  cursor = db.cursor()
  cursor.execute("CREATE TABLE IF NOT EXISTS iscritti (id INTEGER, nome TEXT, tipo TEXT, crazy TEXT, pettorale INTEGER, datanascita DATE, sesso TEXT, pranzo TEXT)")
  cursor.execute("CREATE TABLE IF NOT EXISTS tempi (pettorale INTEGER, tempo TEXT, ntempo TEXT, ftempo DECIMAL)")
  cursor.execute("CREATE VIEW IF NOT EXISTS arrivi AS SELECT i.nome, i.pettorale, i.tipo, i.crazy, i.datanascita, i.sesso, t.tempo, t.ntempo, t.ftempo FROM iscritti i, tempi t WHERE i.pettorale = t.pettorale")
  
  cursor.execute("CREATE VIEW IF NOT EXISTS c_assoluti AS SELECT a.* FROM arrivi AS a ORDER BY a.ftempo ASC")
  cursor.execute("CREATE VIEW IF NOT EXISTS c_crazy AS SELECT a.* FROM arrivi AS a WHERE a.crazy = \"SI\" ORDER BY a.ftempo ASC")
  cursor.execute("CREATE VIEW IF NOT EXISTS c_uomini AS SELECT a.* FROM arrivi AS a WHERE a.sesso = \"M\" ORDER BY a.ftempo ASC")
  cursor.execute("CREATE VIEW IF NOT EXISTS c_donne AS SELECT a.* FROM arrivi AS a WHERE a.sesso = \"F\" ORDER BY a.ftempo ASC")
  cursor.execute("CREATE VIEW IF NOT EXISTS c_ciclisti AS SELECT a.* FROM arrivi AS a WHERE a.tipo = \"CICLISTA\" ORDER BY a.ftempo ASC")
  cursor.execute("CREATE VIEW IF NOT EXISTS c_podisti AS SELECT a.* FROM arrivi AS a WHERE a.tipo = \"PODISTA\" ORDER BY a.ftempo ASC")
  cursor.execute("CREATE VIEW IF NOT EXISTS c_eta AS SELECT * FROM c_crazy ORDER BY datanascita")
  
  
  db.commit()


if __name__ == "__main__":
    main()

